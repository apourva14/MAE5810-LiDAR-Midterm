"""Build the course report from controlled studies and saved scientific figures.

Run after generating the saved sensor and navigation studies:
    python tools/build_report.py --output-dir outputs/reports
All defaults are relative to this project, so the folder can be copied elsewhere.
"""
import argparse
from pathlib import Path
from copy import deepcopy
import json

PROJECT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output-dir', type=Path, default=PROJECT/'outputs/reports',
                    help='Folder for the editable Word report')
parser.add_argument('--data-dir', type=Path, default=PROJECT/'data',
                    help='Root containing the saved evaluation and navigation studies')
parser.add_argument('--figures-dir', type=Path, default=PROJECT/'outputs',
                    help='Root for report figures rendered from saved data')
parser.add_argument('--use-saved-figures', action='store_true',
                    help='Require existing report PNGs and do not rerender figures')
args = parser.parse_args()
DATA = args.data_dir.resolve()
OUTPUTS = args.figures_dir.resolve()
FIGS = OUTPUTS/'evaluation'
EVAL_DATA = DATA/'evaluation'
OUT = args.output_dir.resolve()/'Midterm_Project_Report.docx'
required = ('evaluation/geometry/summary.json', 'evaluation/response/summary.json',
            'navigation_dependencies/summary.json', 'navigation_threshold/summary.json')
missing = [str(DATA/name) for name in required if not (DATA/name).is_file()]
if missing:
    parser.error('Generate the required saved studies first. Missing: '+', '.join(missing))
if args.use_saved_figures:
    required_figures = (
        'weather_report/weather_fov_los.png',
        'weather_report/fov_los.json',
        'all_case_report/target_characteristics.png',
        'all_case_report/weather_measurements.png',
        'all_case_report/weather_trajectories.png',
        'all_case_report/weather_profiles.png',
        'all_case_report/weather_maps.png',
        'navigation_validation/report_threshold_detection.png',
    )
    missing = [str(OUTPUTS/name) for name in required_figures if not (OUTPUTS/name).is_file()]
    if missing:
        parser.error('Render the required report figures first. Missing: '+', '.join(missing))

import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.45)
sec.different_first_page_header_footer = True

def font(style, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for attr in ['ascii','hAnsi','eastAsia','cs']:
        rf.set(qn('w:'+attr),'Times New Roman')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
        rf.attrib.pop(qn('w:'+attr),None)

for name in ['Normal','Body Text','Caption','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s = doc.styles[name]
    font(s, 11)
    s.paragraph_format.line_spacing = 1.5
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after = Pt(5)
font(doc.styles['Title'],17,True)
font(doc.styles['Heading 1'],12,True)
font(doc.styles['Heading 2'],11,True)
font(doc.styles['Heading 3'],11,True)
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles[name].paragraph_format.space_before = Pt(7)
    doc.styles[name].paragraph_format.space_after = Pt(4)
font(doc.styles['Caption'],11)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.space_after = Pt(7)
doc.styles['Caption'].paragraph_format.line_spacing = 1.5
doc.styles['Normal'].paragraph_format.widow_control = True

# Use explicit black styles and no inherited title borders.
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style = doc.styles[name]
    for el in style.element.xpath('.//w:pBdr'):
        el.getparent().remove(el)
    for el in style.element.xpath('.//w:color'):
        for attr in ['themeColor','themeTint','themeShade']:
            el.attrib.pop(qn('w:'+attr),None)

settings = doc.settings.element
mp = OxmlElement('m:mathPr')
mf = OxmlElement('m:mathFont'); mf.set(qn('m:val'),'Cambria Math'); mp.append(mf)
settings.append(mp)

needs_page=False
def p(text='', style=None, align=None, after=None):
    global needs_page
    para = doc.add_paragraph(text, style)
    if needs_page:
        para.paragraph_format.page_break_before=True
        needs_page=False
    if align is not None: para.alignment=align
    if after is not None: para.paragraph_format.space_after=Pt(after)
    return para

def h(text,level=1):
    return p(text, 'Heading '+str(level))

def newpage():
    global needs_page
    needs_page=True

def blank(inches):
    para=p()
    para.paragraph_format.line_spacing=1
    para.paragraph_format.space_after=Pt(0)
    para.paragraph_format.space_before=Pt(0)
    para.paragraph_format.line_spacing=Pt(inches*72)
    para.add_run(' ')
    return para

# Office Math constructs keep the equations editable in Word.
def mr(text,normal=False):
    el=OxmlElement('m:r')
    if normal:
        prop=OxmlElement('m:rPr'); sty=OxmlElement('m:sty'); sty.set(qn('m:val'),'p'); prop.append(sty); el.append(prop)
    prop=OxmlElement('w:rPr'); rf=OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'),'Cambria Math'); rf.set(qn('w:hAnsi'),'Cambria Math')
    prop.append(rf); el.append(prop)
    t=OxmlElement('m:t'); t.set(qn('xml:space'),'preserve'); t.text=str(text); el.append(t)
    return el

def put(parent,items):
    if not isinstance(items,(list,tuple)): items=[items]
    for item in items:
        parent.append(mr(item) if isinstance(item,(str,int,float)) else deepcopy(item))

def sub(base,index):
    el=OxmlElement('m:sSub'); e=OxmlElement('m:e'); put(e,base); el.append(e)
    s=OxmlElement('m:sub'); put(s,index); el.append(s); return el

def sup(base,index):
    el=OxmlElement('m:sSup'); e=OxmlElement('m:e'); put(e,base); el.append(e)
    s=OxmlElement('m:sup'); put(s,index); el.append(s); return el

def frac(num,den):
    el=OxmlElement('m:f')
    n=OxmlElement('m:num'); put(n,num); el.append(n)
    d=OxmlElement('m:den'); put(d,den); el.append(d); return el

def root(expr):
    el=OxmlElement('m:rad'); pr=OxmlElement('m:radPr')
    hide=OxmlElement('m:degHide'); hide.set(qn('m:val'),'1'); pr.append(hide); el.append(pr)
    el.append(OxmlElement('m:deg')); e=OxmlElement('m:e'); put(e,expr); el.append(e); return el

eqno=0
def eq(*pieces):
    global eqno
    eqno+=1
    para=doc.add_paragraph()
    pf=para.paragraph_format
    pf.line_spacing=1.5; pf.space_before=Pt(2); pf.space_after=Pt(4)
    pf.keep_together=True
    pf.tab_stops.add_tab_stop(Inches(3.04),WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    para.add_run('\t')
    math=OxmlElement('m:oMath'); put(math,pieces); para._p.append(math)
    para.add_run('\t('+str(eqno)+')')
    return para

def figure(name, max_h):
    path=FIGS/name
    with Image.open(path) as im: w,h0=im.size
    scale=min(6.5/w,max_h/h0)
    para=p(align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
    para.paragraph_format.line_spacing=1
    para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(path),width=Inches(w*scale),height=Inches(h0*scale))
    return para

def caption(text):
    return p(text,'Caption')

def table(headers,rows,widths,spacing=1.5,size=None,pad_y=65):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        el=OxmlElement('w:'+side); el.set(qn('w:w'),str(pad_y) if side in ['top','bottom'] else '90'); el.set(qn('w:type'),'dxa'); margins.append(el)
    pr.append(margins)
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for rowidx, values in enumerate([headers]+rows):
        row=t.rows[0] if rowidx==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if rowidx==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for i,value in enumerate(values):
            cell=row.cells[i]; cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para=cell.paragraphs[0]; para.paragraph_format.line_spacing=spacing
            para.paragraph_format.space_after=Pt(0)
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
            run=para.add_run(value); run.bold=(rowidx==0)
            if size is not None: run.font.size=Pt(size)
            if rowidx==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EEEEEE'); cell._tc.get_or_add_tcPr().append(shade)
    p(after=1).paragraph_format.line_spacing=Pt(1)
    return t

foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
font(doc.styles['Footer'],11)
doc.core_properties.title='Physics Based Simulation of a Three Dimensional LiDAR Sensor'
doc.core_properties.subject='MAE 5810 Robot Perception Midterm Project'
doc.core_properties.author=''
doc.core_properties.keywords='LiDAR, physics based model, FOV, LOS, probabilistic measurement'

# Report source uses actual saved study outputs; no simulation occurs here.
import csv
def records(name):
    with (EVAL_DATA/name).open(newline='') as stream:
        return list(csv.DictReader(stream))

G = json.loads((EVAL_DATA/'geometry/summary.json').read_text())
R = json.loads((EVAL_DATA/'response/summary.json').read_text())
NAV_DATA = DATA/'navigation_dependencies'
NAV_FIGS = OUTPUTS/'navigation_dependencies'
WEATHER_FIGS = OUTPUTS/'all_case_report'
WEATHER_COVERAGE = OUTPUTS/'weather_report'
N = json.loads((NAV_DATA/'summary.json').read_text())
NC = {item['id']: item for item in N['cases']}
NB = NC['baseline']
distance_examples=[]
for case_id in ('target_near','baseline','target_far'):
    with (NAV_DATA/NC[case_id]['representative_replay']/'traces.csv').open(newline='') as stream:
        row=next(row for row in csv.DictReader(stream) if abs(float(row['t'])-6.0)<1e-9)
    distance_examples.append((float(row['eligible_range_sum_m'])/int(row['eligible_echoes']),
                              1000*float(row['eligible_power_sum_w'])/int(row['eligible_echoes'])))
import sys
sys.path.insert(0,str(PROJECT))
if not args.use_saved_figures:
    from lidar_project.visualization.all_case_report import render_weather_report
    render_weather_report(NAV_DATA, WEATHER_FIGS)
VALIDATION_FIGS = OUTPUTS/'navigation_validation'
T = json.loads((DATA/'navigation_threshold/summary.json').read_text())
TA = T['aggregate']
if not args.use_saved_figures:
    from lidar_project.visualization.navigation_validation import _threshold
    from lidar_project.visualization import theme
    import matplotlib.pyplot as plt
    VALIDATION_FIGS.mkdir(parents=True, exist_ok=True)
    with plt.rc_context({**theme.MATPLOTLIB, "font.size": 8.5}):
        _threshold(T, VALIDATION_FIGS/'report_threshold_detection.png')
    from lidar_project.visualization.weather_coverage import prepare_weather_coverage, render_weather_coverage
    render_weather_coverage(prepare_weather_coverage(NAV_DATA, WEATHER_COVERAGE), WEATHER_COVERAGE)
C = json.loads((WEATHER_COVERAGE/'fov_los.json').read_text())

# 1 -- Title page.
blank(1.1)
p('Physics Based Simulation of a\nThree Dimensional LiDAR Sensor','Title',WD_ALIGN_PARAGRAPH.CENTER,18)
p('MAE 5810 Robot Perception\nMidterm Project Report','Subtitle',WD_ALIGN_PARAGRAPH.CENTER,24)
p('Author  __________________________________','Normal',WD_ALIGN_PARAGRAPH.CENTER,18)
p('Student ID  _______________________________','Normal',WD_ALIGN_PARAGRAPH.CENTER,24)
p('Instructor  Silvia Ferrari','Normal',WD_ALIGN_PARAGRAPH.CENTER,12)
p('October 19 2026','Normal',WD_ALIGN_PARAGRAPH.CENTER)
newpage()

# 2 -- Motivation and nominal scan model.
h('Abstract')
p(f'This project models a spinning three dimensional pulsed LiDAR with opaque visibility, surface response, attenuation, and receiver noise. Evaluation has two parts: nine target-characteristic settings at matched robot poses, and clear air, rain, and fog navigation in one fixed obstacle setup. Nine recorded weather runs retain all trajectories and measurement profiles. Attenuation reduces mean observed endpoint cells from 73.7 to 47.1 percent while all nine runs reach the goal. An independent 33,000-trial experiment checks detection near threshold; probability calibration RMS error is {100*TA["calibration_rmse"]:.2f} percentage points. These results establish internal model consistency and conditional environmental dependencies. Three navigation seeds are descriptive; physical calibration and general reliability remain untested.')
h('1 Introduction')
p('A virtual LiDAR must determine where its beams can reach and whether their returned signal is measurable. A missing point may arise from angular sampling, occlusion, a range gate, or weak power. This project therefore models measurements together with field of view (FOV) and line of sight (LOS), addressing the dependencies on sensor state, targets, and environment required by the midterm assignment [1].')
p('The navigation sensor uses eight channels. Surface modeling is motivated by Heiden et al. [2], whose continuous wave phase estimator differs from this pulsed geometric approximation. Target experiments quantify geometric exposure and surface response; weather experiments hold the physical scene, robot, sensor, reference route, and controller fixed. Threshold trials resolve uncertain detections independently of navigation.')
h('2 Theory and Methods')
h('2.1 Scan pattern',2)
p('The sensor frame has x forward, y left, and z up. World z is up. With M azimuth samples per full revolution and L channels, beam angles and the nominal unit direction are')
eq(sub('α','i'),' = ',sub('α','min'),' + ',frac('2πi','M'),',     ',sub('β','j'),' = ',sub('β','min'),' + ',frac(['j(',sub('β','max'),' − ',sub('β','min'),')'],'L − 1'))
eq(sub('b','ij'),' = ',sup(['[cos ',sub('β','j'),' cos ',sub('α','i'),', cos ',sub('β','j'),' sin ',sub('α','i'),', sin ',sub('β','j'),']'],'T'))
p('A full sweep uses i = 0, …, M − 1. A partial window uses the same azimuth spacing. A single channel takes the configured upper elevation. Angles in equations are in radians; plots use degrees.',after=0)
newpage()

# 3 -- Geometry and physics model.
h('2.2 Pose and coverage',2)
p('All channels at one azimuth share a firing time. The origin translates at constant world velocity; yaw may change during a scan while roll and pitch remain fixed:')
eq(sub('t','i'),' = ',frac([sub('α','i'),' − ',sub('α','min')],['2π',sub('f','rot')]),',     ',sub('o','i'),' = ',sub('o','0'),' + v',sub('t','i'))
p('The sensor to world rotation is R = Rz(yaw) Ry(pitch) Rx(roll), and the world beam is d = Rb. For point p, let q = Rᵀ(p − o), r = ‖q‖, a = atan2(qy, qx), and e = arcsin(qz/r). The nominal FOV is')
eq('F = {p : ',sub('r','min'),' ≤ r ≤ ',sub('r','max'),', a ∈ A, ',sub('β','min'),' ≤ e ≤ ',sub('β','max'),'}')
p('A is the circular azimuth window. Discrete beams can miss targets within this continuous set. For each subray, the nearest opaque intersection with scene surfaces S is')
eq(sub('r','k'),' = min {s > ',sup('10','−6'),' m : o + s',sub('d','k'),' ∈ S}')
p('The first hit is found before range gating, so a weak or too near object still blocks farther geometry. A finite beam uses normalized Gaussian weighted subrays in the plane normal to its central direction. For full divergence δ, the angular standard deviation is δ/4; a sunflower pattern truncated at 2.5 standard deviations supplies the sample offsets.')
h('2.3 Surface power and echo formation',2)
p('For a ray facing normal n, μ = max(0, −nᵀd). Diffuse weight ρd, specular weight ρs, and exponent m define the approximate surface law')
eq('f = ',frac(sub('ρ','d'),'π'),' + ',frac([sub('ρ','s'),'(m + 2)'],'2π'),sup(['[max(2',sup('μ','2'),' − 1, 0)]'],'m'))
eq(sub('P','k'),' = ',sub('w','k'),sub('P','t'),sub('A','r'),'η f',frac(sub('μ','k'),sup(sub('r','k'),'2')),' exp(−2κ',sub('r','k'),')')
p('Here Ar = π(Dr/2)², η is optical efficiency, and κ = κair + arain qrainᵇ. Invalid hits contribute zero power. Rain and fog change attenuation, without backscatter returns. Consecutive subray hits separated by at most the pulse gate merge into an echo group G:')
eq(sub('Δr','gate'),' = ',frac(['c',sub('τ','FWHM')],'2'),',     ',sub('P','G'),' = Σ',sub('P','k'),',     ',sub('r̄','G'),' = ',frac(['Σ',sub('P','k'),sub('r','k')],['Σ',sub('P','k')]),'     (k ∈ G)')
p('The centroid approximates pulsed ranging rather than waveform processing; it can lie between surfaces at edges. Geometry supports planes, spheres, yawed boxes, and capped cylinders.',after=0)
newpage()

# 4 -- Uncertainty and graduate mathematical analysis.
h('2.4 Detection and range uncertainty',2)
p('For an existing echo, noisy power is P̃ = PG + ν, with ν ∼ N(0, σP²). Detection requires P̃ ≥ γσP, giving the single echo probability')
eq(sub('P','D'),'(G) = Φ(',frac(sub('P','G'),sub('σ','P')),' − γ)')
p('Φ is the standard normal cumulative distribution function. The strongest mode selects the greatest noisy detected power; first and last modes select by range. There is at most one retained return per nominal beam. With S = P̃/σP, the assumed timing inspired range law is')
eq(sup(sub('σ','r'),'2'),'(S) = ',sup(sub('σ','0'),'2'),' + ',frac(sup(['(c',sub('t','rise'),'/2)'],'2'),'2 max(S, 10⁻⁹)'),',     z = ',sub('r̄','G'),' + ',sub('σ','r'),'(S)ξ')
p('ξ is independent standard normal noise. This law is an assumption, not a derived Cramér–Rao bound. Selection truncates measured power but leaves the subsequent standardized range noise normal. Coordinates use the central beam and known firing pose:')
eq(sup('p̂','S'),' = z b,     ',sup('p̂','W'),' = o + z d')
h('2.5 Mathematical dependencies',2)
p('A forward plane intersection is r = nᵀ(p0 − o)/(nᵀd). Horizontal ground at mounting height h and downward world beam elevation β therefore gives')
eq('r = −',frac('h','sin β'),',     ',frac('∂r','∂h'),' = −',frac('1','sin β'),',     ',frac('∂r','∂β'),' = ',frac('h cos β',[sup('sin','2'),' β']))
p('Height and orientation change range, with increasing sensitivity near horizontal beams. The FOV transformation changes coverage with pose; the nearest hit rule changes LOS at occluder edges. For a beam filling diffuse target, P = C exp(−2κr)/r² and C = PtArηρd cos θ/π, so')
eq(frac('∂ ln P','∂r'),' = −',frac('2','r'),' − 2κ,     ',frac('∂ ln P','∂κ'),' = −2r,     ',frac('∂ ln P','∂θ'),' = −tan θ')
p('Diffuse weight scales power linearly. Greater range, extinction, and incidence angle weaken the diffuse signal, and therefore reduce detection probability. The angle derivative holds range fixed and excludes the specular branch. For any selected physical reference,')
eq('E[',sup(['(z − ',sub('r','ref'),')'],'2'),' | P̃, G] = ',sup(['(',sub('r̄','G'),' − ',sub('r','ref'),')'],'2'),' + ',sup(sub('σ','r'),'2'),'(S)')
p('The first term is centroid displacement; the second is conditional random variance. Angular, pose, calibration, and additional physical model errors are excluded.',after=0)
newpage()

# 5 -- Reproducible two-part evaluation.
h('3 Procedure')
p('Python and NumPy generate measurements; independent renderers read saved results. The fixed navigation scene has 12 objects: nine in the planning map and three unknown to the planner. Its designated target is a pure diffuse box centered at (11, 1.1, 0.6) m, sized (0.8, 1.6, 1.2) m, with diffuse weight 0.30.')
p('The sensor uses eight channels from -20 to +10 degrees, 360 azimuths, one subray, a level 0.6 m mount, and 0.5-200 m gates. Parameters are 75 W peak power, 25 mm aperture, 0.5 efficiency, 4 × 10⁻⁹ W power noise RMS, threshold factor 10, 0.01 m range floor, 5 ns pulse, and 1 ns rise time.')
p('A* supplies a shared reference from (0, 0) to (18, 0) m on a 0.25 m grid, with 0.90 m total inflation. The measured-point controller uses 0.12 ≤ z ≤ 1 m without object labels. Limits are 1.2 m/s, 90 degrees/s, 0.2 s steps, and 36 s. Forward Euler integrates')
eq(sub('p','k+1'),' = ',sub('p','k'),' + ',sub('v','k'),'Δt ',sup(['[cos ',sub('θ','k'),', sin ',sub('θ','k'),']'],'T'),',     ',sub('θ','k+1'),' = wrap(',sub('θ','k'),' + ',sub('ω','k'),'Δt)')
h('3.1 Target characteristics',2)
p('Nine target settings vary diffuse weight (0.02, 0.30, 0.85), yaw (0, 45, 75 degrees), height (0.3, 1.2, 2.0 m), and center x (9, 11, 13 m). Each is replayed at the same 271 baseline snapshots across three seeds. The shared baseline appears in both evaluation parts. Count target intersections, nearest visible hits, in-gate eligible echoes NE, and detections ND separately. For qi from Equation (9),')
eq(sub('p̂','D'),' = ',frac(sub('N','D'),sub('N','E')),',     ',sub('p̄','D'),' = ',frac(['Σ',sub('q','i')],sub('N','E')),'     (i ∈ E)')
h('3.2 Environmental conditions through navigation',2)
p('Hold all obstacle geometry, target properties, sensor settings, robot parameters, controller, and reference route fixed. Compare clear air (κ = 0.0002 m⁻¹), rain at 50 mm/h (effective κ = 0.104764 m⁻¹), and fog represented by κ = 0.6 m⁻¹. Each condition uses seeds 1-3: nine closed-loop runs and nine matched-pose replays, including the shared baseline. All recorded trajectories and metric histories are retained.')
p('Points per scan include all detected surfaces. Mapping reports endpoint-observed cells, not free space or accuracy. Target range RMS pools squared errors over detected replay returns; misses are not zero errors. Weather coefficients are illustrative attenuation settings; no backscatter or false weather returns are modeled. Three navigation seeds provide descriptive comparisons, not a reliability estimate.',after=0)
newpage()

# 6 -- Target-characteristic evaluation and unchanged-scene visibility.
h('4 Results')
h('4.1 Target characteristics',2)
figure(WEATHER_COVERAGE/'weather_fov_los.png',2.35)
caption('Figure 1  Unchanged baseline at 2.6 s: 2,880 emitted rays; 12 visible and 12 existing-wall-blocked target rays. Only representative rays and relevant true surfaces are drawn. The full-circle angular envelope is clipped at 4 m; target rays extend farther. This is explanatory geometry, not measured map points.')
figure(WEATHER_FIGS/'target_characteristics.png',3.05)
caption('Figure 2  All nine target settings at matched baseline poses. Power is the noiseless eligible-echo mean; counts distinguish intersections, eligibility, and detections. Range RMS is conditional on detected target returns. Target position, orientation, and size also change geometric exposure.')
p('Diffuse weights 0.02, 0.30, and 0.85 give power ratios 0.0667, 1, and 2.8333. At one common seed-1 viewpoint (6 s), target centers x = 9, 11, and 13 m give ranges 2.134, 4.011, and 5.950 m and powers 0.339830, 0.105144, and 0.048696 mW. Whole-route averages need not preserve this ordering because eligible exposures and incidence differ.',after=0)
newpage()

# 7 -- Environmental conditions: scalar metrics and matched-pose profiles.
h('4.2 Environmental measurement metrics',2)
caption('Table 1  Same fixed scene; three seeds per weather condition. Navigation/map values are means of run summaries. Range RMS and detection fraction pool target echoes at identical baseline poses.')
weather_ids=('baseline','rain','fog')
weather_names={'baseline':'Clear','rain':'Rain 50 mm/h','fog':'Fog κ = 0.6 m⁻¹'}
table(['Condition','Mean points / scan','Observed cells (%)','Target RMS (mm)','Detected / eligible (%)'],[
 [weather_names[c], f"{NC[c]['aggregate']['mean_returns_per_control_step']:,.1f}",
  f"{100*NC[c]['aggregate']['mean_observed_fraction']:.2f}",
  f"{1000*NC[c]['replay']['range_rmse_m']:.3f}",
  f"{100*NC[c]['replay']['observed_detection_probability']:.2f}"] for c in weather_ids
], [1.5,1.2,1.2,1.15,1.45],spacing=1.0,size=11,pad_y=65)
figure(WEATHER_FIGS/'weather_measurements.png',2.6)
caption('Figure 3  Weather measurement profiles over the same saved baseline poses. Colors denote weather; all three seeds contribute. Cumulative range RMS uses pooled detected-return squared errors and its actual detection denominator. No profile is extended beyond its recorded exposure.')
p(f'Rain and fog reduce mean eligible target power from {NB["replay"]["mean_eligible_power_w"]*1000:.6f} to {NC["rain"]["replay"]["mean_eligible_power_w"]*1000:.6f} and {NC["fog"]["replay"]["mean_eligible_power_w"]*1000:.6f} mW. This follows the two-way attenuation term in Equation (7); target geometry and angular sampling are unchanged.')
p('Fog lowers mean points per scan by 20.8 percent and observed endpoint coverage by 26.57 percentage points relative to clear air. Target RMS rises from 9.992 to 11.139 mm, while detection falls to 98.07 percent of eligible target echoes. These describe different denominators: whole-scene closed-loop returns, common-region endpoint cells, and detected-target replay errors.')
p('The range floor dominates strong nearby returns, so a substantial optical-power loss can produce a small RMS increase. Detection and coverage expose losses that conditional range RMS alone can conceal. Missing returns remain in the detection denominator; map gaps are not filled using scene truth.',after=0)
newpage()

# 8 -- Independent near-threshold detection and uncertainty.
h('4.3 Near-threshold detection and range uncertainty',2)
figure(VALIDATION_FIGS/'report_threshold_detection.png',2.2)
caption('Figure 4  Detection and calibration residuals from 33,000 actual scan trials. Each marker counts 1,000 trials, including misses; bars are pointwise 95% Wilson intervals. Small horizontal offsets distinguish viewpoints. The shaded band is the predicted 5-95% transition, not an uncertainty interval.')
p('The weather replays nearly saturate detection. For independent component validation, select three eligible rays from this same baseline scene at ranges 7.99, 4.90, and 1.93 m. Vary extinction to obtain 11 mean-SNR settings from 6 to 14, with 1,000 independent actual scans per setting. Geometry, diffuse target response, receiver noise, and threshold remain fixed. Ray selection does not depend on noisy detection; these pencil-ray probes are not additional navigation runs.')
p(f'At mean SNR 10, predicted probability is 0.5; measured fractions are 0.510, 0.492, and 0.498. Over 33 settings, calibration RMS error is {100*TA["calibration_rmse"]:.3f} percentage points, with maximum residual {100*TA["max_abs_detection_residual"]:.3f} points. The model lies within {TA["pointwise_wilson_contains_model"]}/33 pointwise intervals; the excluded prediction is retained.')
p('With n independent trials, observed fraction p̂, and a = 1.96, the Wilson interval is')
eq(frac(['p̂ + ',frac(sup('a','2'),'2n'),' ± a ',root([frac('p̂(1 − p̂)','n'),' + ',frac(sup('a','2'),['4',sup('n','2')])])],['1 + ',frac(sup('a','2'),'n')]))
p('For detected echoes only, define range and standardized errors; misses are not zero errors:')
eq(sub('e','i'),' = ',sub('z','i'),' − ',sub('r','true,i'),',     ',sub('ε','i'),' = ',frac(sub('e','i'),[sub('σ','r'),'(',sub('S','i'),')']),',     RMSE = ',root(frac(['Σ',sup(sub('e','i'),'2')],sub('N','D'))))
p(f'Across {TA["detections"]:,} detections, range RMS is {1000*TA["conditional_range_rmse_m"]:.3f} mm versus {1000*TA["predicted_conditional_range_rmse_m"]:.3f} mm predicted from measured-SNR conditional variance; standardized RMS is {TA["normalized_range_rmse"]:.4f}. Agreement checks the assumed noise and selection laws, not physical accuracy. Intervals do not provide simultaneous coverage.',after=0)
newpage()

# 9 -- All environmental trajectories and recorded metric histories.
h('4.4 Environmental conditions through navigation',2)
figure(WEATHER_FIGS/'weather_trajectories.png',3.0)
figure(WEATHER_FIGS/'weather_profiles.png',2.5)
caption('Figure 5  All nine weather runs in the unchanged scene. Upper: recorded 3D paths and top-down detail. Lower: points per scan, endpoint coverage, swept clearance, and route error. Exact overlaps remain overlaps. Lines stop at recorded termination; terminal poses add no scan. Colors denote weather, line styles seeds.')
caption('Table 2  Three runs per condition; all reach the goal without collision or timeout. Gap is the worst swept clearance. Three seeds are descriptive.')
table(['Condition','Goals','Mean time (s)','Mean path (m)','Min gap (m)'],[
 [weather_names[c],f"{NC[c]['aggregate']['goals']}/3",f"{NC[c]['aggregate']['mean_elapsed_s']:.2f}",
  f"{NC[c]['aggregate']['mean_path_length_m']:.3f}",f"{NC[c]['aggregate']['worst_surface_gap_m']:.3f}"] for c in weather_ids
], [1.7,.6,1.35,1.4,1.45],spacing=1.0,size=11,pad_y=40)
newpage()

# 10 -- Measured maps, effectiveness, conclusions and references.
h('4.5 Route generation and measured 3D mapping',2)
figure(WEATHER_FIGS/'weather_maps.png',2.2)
caption('Figure 6  Seed-1 measured maps for clear air, rain, and fog in the same fixed scene. Colors show height; all panels use the same axes and height scale. Display hides z ≤ 0.12 m. Each map ends with its own run; no true scene surfaces fill gaps.')
p('Measured surfaces accumulate in 0.12 m voxels. Endpoint cells are 0.2 m over [-2, 21] m × [-6, 7] m (NROI = 7,475). For K intervals and K + 1 states,')
eq('L = Σ ‖',sub('p','k+1'),' − ',sub('p','k'),'‖,     ',sub('f','obs'),' = ',frac(sub('N','obs'),sub('N','ROI')),',     k = 0, …, K − 1')
p('Mean observed endpoint cells are 73.67 percent in clear air, 70.33 in rain, and 47.10 in fog. All weather runs reach the goal; endpoint observations do not establish accuracy or free space.')
p(f'Independent checks cover {G["analytic_cases"]} geometry and {R["radiometry"]["rows"]} power cases; maximum residuals are {G["max_abs_range_error_m"]:.2e} m and {R["radiometry"]["max_abs_relative_power_error"]:.2e} relative power.')
h('5 Conclusions')
p('Target replays expose surface response and geometric sampling. With obstacles fixed, attenuation reduces returns and coverage and increases conditional range error, although all nine weather runs succeed. Threshold trials support the assumed noise model. Three seeds do not establish navigation reliability. Physical calibration, backscatter, pose uncertainty, moving objects, and scan motion remain future work.')
h('References')
p('[1] “Midterm project,” Robot Perception, course assignment handout, pp. 1-2.').paragraph_format.line_spacing=1.0
p('[2] E. Heiden, Z. Liu, R. K. Ramachandran, and G. S. Sukhatme, “Physics-based Simulation of Continuous-Wave LIDAR for Localization, Calibration and Tracking,” Proc. IEEE ICRA, 2020, pp. 2595-2601.').paragraph_format.line_spacing=1.0

for shape,description in zip(doc.inline_shapes,[
    'Full-circle field of view and natural target occlusion by an existing wall in the unchanged weather scene.',
    'Nine matched-pose target-characteristic settings with eligible optical power, geometric sampling, detected-return range RMS and target features.',
    'Clear, rain and fog measurement profiles at the same baseline poses, including detected target returns, power and pooled cumulative range RMS.',
    'Three near-threshold target-ray detection curves and calibration residuals with pointwise Wilson intervals.',
    'All nine recorded weather trajectories in the unchanged three dimensional obstacle scene with a top-down detail.',
    'Recorded weather histories of returns, endpoint coverage, swept clearance and spatial route error, retaining all seeds and true termination times.',
    'Measured three dimensional endpoint maps in clear air, rain and fog, sharing physical axes and height colors.'
]):
    shape._inline.docPr.set('descr',description)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(json.dumps({'docx':str(OUT),'equations':eqno,'figures':len(doc.inline_shapes)}))
